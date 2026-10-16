"""Reads the brand's knowledge spreadsheet (Products, Routines, Rules) into JSON:
a plain transcription, no rule applied. Run by `npm run advisor:import`, which
applies the rules (scripts/lib/advisory/product-knowledge.mjs,
routine-knowledge.mjs) and writes the file into git, so every correction made in
the spreadsheet is a reviewable diff (docs/storefront-advisor-plan.md § 2).

Each tab's header row is the one whose cells end in « (key) »; a row is kept when
its first key column is filled. Each row carries `_row`, its line in the
spreadsheet, so a rejected row can be found again.

    python scripts/advisor/export-knowledge-sheet.py <file.xlsx>   → JSON on stdout
"""
import json, re, sys

import openpyxl

KEY = re.compile(r'\(([a-z_]+)\)\s*$')
# The first column a row needs to count, per tab (a title or section row has none).
REQUIRED = {'Products': 'handle', 'Routines': 'functional_type', 'Rules': 'key'}


def value(v):
    if isinstance(v, bool):
        return 'Oui' if v else 'Non'
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def tab(ws, required):
    for r in range(1, min(ws.max_row, 30) + 1):
        keys = [KEY.search(str(c.value or '')) for c in ws[r]]
        if sum(1 for k in keys if k) >= 3:
            break
    else:
        sys.exit(f'{ws.title}: no header row with « (key) » names')
    cols = [(i, k.group(1)) for i, k in enumerate(keys) if k]
    rows = []
    for row in ws.iter_rows(min_row=r + 1):
        record = {key: value(row[i].value) if i < len(row) else None for i, key in cols}
        if record.get(required) is None:
            continue
        rows.append({'_row': row[0].row, **record})
    return rows


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {name: tab(wb[name], required) for name, required in REQUIRED.items()}
    json.dump({k.lower(): v for k, v in out.items()}, sys.stdout, ensure_ascii=False, indent=1)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main(sys.argv[1])
